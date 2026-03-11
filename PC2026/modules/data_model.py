"""
Modelo de datos: carga, procesa y expone DataFrames listos para el dashboard.

Flujo:
1. Descarga Excel de Google Drive
2. Lee hoja "Total Cuentas Gamma" como master
3. Lee hojas individuales de BDM para extraer colores de heatmap
4. Fusiona todo en un modelo unificado
"""
from __future__ import annotations

from dataclasses import dataclass, field

import pandas as pd
import streamlit as st

from modules.excel_processor import (
    SECTIONS,
    SKIP_SHEETS,
    SheetData,
    read_worksheet_with_colors,
)
from modules.google_drive import download_excel_from_drive
from modules.scoring import (
    build_score_summary,
    detect_unique_colors,
    score_color,
    score_color_dataframe,
    score_color_labels,
)

# Columnas clave que se auto-detectan
KEY_COLUMNS = {
    "cliente": ["cliente", "client"],
    "comercial": ["comercial", "ejecutivo", "bdm", "owner"],
    "zona": ["zona", "region"],
    "sector": ["sector", "industria"],
    "tipo_cuenta": ["tipo de cuenta", "tipo_cuenta", "tipo cuenta"],
    "nivel_relacion": ["nivel de relacionamiento", "nivel_relacion", "relacion"],
    "valoracion": ["valorizacion de la relación", "valorizacion", "valoracion"],
    "estrategia": ["estrategia"],
    "pasos": ["pasos a seguir", "pasos"],
    "descripcion": ["descripción general", "descripcion"],
}

# Columnas de heatmap (productos + servicios)
HEATMAP_PRODUCTS = [
    "CYBERACADEMY", "EDR", "ANTIMALWARE", "DLP", "MIGROSEGMENTACIÓN", "NGFW",
    "Balanceo de carga GSLB", "SWITCHES", "WIRELESS", "SD-WAN", "ZTNA", "NAC",
    "SEGURIDAD EN CORREO (SEG)", "PROTECCION DE APLICACIONES WEB (WAF/WAAP)",
    "BALANCEO DE CARGA (ADC)", "AntiDDoS",
    "MULTIPLE FACTOR DE AUTENTICACIÓN (MFA)", "SEGURIDAD PARA ACCESO CLOUD (CASB)",
    "SASE", "GESTION DE ACCESO PRIVILEGIADO (PAM)", "PROTECCIÓN BASES DE DATOS (DBF)",
    "NDR", "ITDR", "SOAR", "SIEM", "SANDBOX", "DECEPTION", "XDR", "GRC", "CNAAP",
    "CLASIFICACION Y ETIQUETADO", "OBSERVABILIDAD", "HIPERCONVERGENCIA",
    "ALMACENAMIENTO", "BACKUP", "CIBERSEGURIDAD DE LA IA", "CRIPTOGRAFIA POST-CUANTICA",
]

HEATMAP_SERVICES = [
    "MONITOREO DE AMENAZAS (CSOC)", "ANALISIS DE ULNERABILIDADES",
    "ETHICAL HACKING", "MONITOREO DE MARCA", "RESPUESTA A INCIDENTES",
    "CONSULTORIA", "PHISHING",
]

RELATIONSHIP_COLS = [
    "Nivel de Relacionamiento", "Valorizacion de la relación",
    "C-level", "Dirección", "Gerente", "Operativo", "Influenciador",
]

AREAS_INTERNAS = [
    "CIBERSEGURIDAD", "SEGURIDAD DE LA INFORMACION", "INFRAESTRUCTURA",
    "RIESGOS", "TRANSFORMACION DIGITAL", "VP", "COMPRAS", "FINANCIERO",
    "JURIDICO", "MERCADEO",
]


def _find_column(headers: list[str], aliases: list[str]) -> str | None:
    headers_lower = {h.lower().strip(): h for h in headers}
    for alias in aliases:
        for h_low, h_orig in headers_lower.items():
            if alias in h_low:
                return h_orig
    return None


def _find_matching_col(df_cols: list[str], target: str) -> str | None:
    """Busca una columna por coincidencia parcial case-insensitive."""
    target_low = target.lower().strip()
    for c in df_cols:
        if target_low in c.lower().strip() or c.lower().strip() in target_low:
            return c
    return None


def _normalize_sector(s) -> str:
    if pd.isna(s) or not s:
        return "Sin Sector"
    s = str(s).strip().upper()
    # Normalizar variantes comunes
    mapping = {
        "FINANCIERO": "Financiero",
        "GOBIERNO": "Gobierno",
        "INDUSTRIA": "Industria",
        "INDUSTRIA Y COMERCIO": "Industria y Comercio",
        "CORPORATIVO": "Corporativo",
        "SALUD": "Salud",
        "EDUCACION": "Educación",
        "EDUCACIÓN": "Educación",
        "RETAIL": "Retail",
        "SERVICIOS": "Servicios",
        "DISTRIBUCIÓN Y COMERCIO": "Distribución y Comercio",
        "DISTRIBUCION Y COMERCIO": "Distribución y Comercio",
    }
    for key, val in mapping.items():
        if key in s:
            return val
    return str(s).strip().title()


def _normalize_comercial(s) -> str:
    if pd.isna(s) or not s:
        return "Sin Asignar"
    return str(s).strip().title()


@dataclass
class ProcessedSheet:
    name: str
    values: pd.DataFrame
    colors: pd.DataFrame
    scores: pd.DataFrame
    labels: pd.DataFrame
    headers: list[str] = field(default_factory=list)
    sections: dict[str, list[str]] = field(default_factory=dict)
    key_cols: dict[str, str] = field(default_factory=dict)

    @property
    def heatmap_product_cols(self) -> list[str]:
        return self.sections.get("heatmap_productos", [])

    @property
    def heatmap_service_cols(self) -> list[str]:
        return self.sections.get("heatmap_servicios", [])

    @property
    def heatmap_all_cols(self) -> list[str]:
        return self.heatmap_product_cols + self.heatmap_service_cols

    @property
    def col_cliente(self) -> str | None:
        return self.key_cols.get("cliente")

    @property
    def col_comercial(self) -> str | None:
        return self.key_cols.get("comercial")


@dataclass
class DashboardData:
    """Datos completos del dashboard."""
    # Master dataframe con toda la info unificada
    master: pd.DataFrame
    # Sheets procesados individualmente (por si se necesitan)
    sheets: dict[str, ProcessedSheet]
    sheet_names: list[str]
    all_sheet_names: list[str]
    excel_bytes: bytes
    unique_colors: dict[str, int]
    # Columnas detectadas
    col_cliente: str
    col_comercial: str
    col_zona: str
    col_sector: str
    col_tipo: str
    # Heatmap columns found in data
    heatmap_product_cols: list[str]
    heatmap_service_cols: list[str]

    def get_sheet(self, name: str) -> ProcessedSheet | None:
        return self.sheets.get(name)

    @property
    def primary_sheet(self) -> ProcessedSheet | None:
        if self.sheet_names:
            return self.sheets.get(self.sheet_names[0])
        return None

    @property
    def bdm_names(self) -> list[str]:
        return sorted(self.master[self.col_comercial].dropna().unique().tolist())

    @property
    def bdm_sheets(self) -> list[str]:
        return [n for n in self.sheet_names if n.lower() != "total cuentas gamma"]

    @property
    def heatmap_all_cols(self) -> list[str]:
        return self.heatmap_product_cols + self.heatmap_service_cols

    @property
    def n_clients(self) -> int:
        return len(self.master)

    @property
    def n_bdms(self) -> int:
        return self.master[self.col_comercial].nunique()

    @property
    def zones(self) -> list[str]:
        return sorted(self.master[self.col_zona].dropna().unique().tolist())

    @property
    def sectors(self) -> list[str]:
        return sorted(self.master[self.col_sector].dropna().unique().tolist())

    @property
    def account_types(self) -> list[str]:
        return sorted(self.master[self.col_tipo].dropna().unique().tolist())


def _build_heatmap_matrix_from_bdm_sheets(
    wb, bdm_sheet_names: list[str]
) -> tuple[pd.DataFrame, dict[str, int]]:
    """
    Lee las hojas individuales de BDM y extrae los colores de heatmap.
    Retorna un DataFrame con scores por producto/servicio para cada fila,
    y un dict de colores únicos.
    """
    all_rows = []
    all_color_rows = []

    for name in bdm_sheet_names:
        ws = wb[name]
        sd = read_worksheet_with_colors(ws)
        if sd.df.empty:
            continue

        # Encontrar columna cliente
        col_cliente = _find_column(sd.raw_headers, KEY_COLUMNS["cliente"])
        if not col_cliente or col_cliente not in sd.df.columns:
            continue

        for idx, row in sd.df.iterrows():
            cliente = row.get(col_cliente)
            if pd.isna(cliente) or not str(cliente).strip():
                continue

            row_data = {"_cliente_key": str(cliente).strip().upper(), "_bdm_sheet": name}
            color_row = {}

            # Extraer colores de heatmap para cada producto/servicio
            for col_name in HEATMAP_PRODUCTS + HEATMAP_SERVICES:
                matched = _find_matching_col(sd.raw_headers, col_name)
                if matched:
                    color_col = f"{matched}_color"
                    if color_col in sd.colors.columns:
                        hex_val = sd.colors.at[idx, color_col] if idx < len(sd.colors) else None
                        sc = score_color(hex_val)
                        row_data[col_name] = sc.score
                        color_row[col_name] = hex_val
                    # Also get text value
                    text_val = row.get(matched, None)
                    row_data[f"{col_name}_text"] = text_val if not pd.isna(text_val) else None

            all_rows.append(row_data)
            all_color_rows.append(color_row)

    if not all_rows:
        return pd.DataFrame(), {}

    scores_df = pd.DataFrame(all_rows)
    colors_df = pd.DataFrame(all_color_rows)

    # Detect unique colors
    unique = detect_unique_colors(colors_df) if not colors_df.empty else {}

    return scores_df, unique


@st.cache_data(ttl=120, show_spinner="Cargando datos desde Google Drive...")
def load_dashboard_data(file_id: str | None = None) -> DashboardData:
    """Pipeline completo: descarga Excel -> procesa -> modelo unificado."""
    from modules.config import SHEET_ID
    import io
    from openpyxl import load_workbook

    fid = file_id or SHEET_ID
    excel_bytes = download_excel_from_drive(fid)

    wb = load_workbook(io.BytesIO(excel_bytes), data_only=True)
    all_sheet_names = wb.sheetnames
    data_sheet_names = [n for n in all_sheet_names if n.lower().strip() not in SKIP_SHEETS]

    sheets: dict[str, ProcessedSheet] = {}
    all_colors = pd.DataFrame()

    try:
        # 1. Procesar todas las hojas
        for name in data_sheet_names:
            ws = wb[name]
            sd: SheetData = read_worksheet_with_colors(ws)

            if sd.df.empty:
                sheets[name] = ProcessedSheet(
                    name=name, values=sd.df, colors=sd.colors,
                    scores=pd.DataFrame(), labels=pd.DataFrame(),
                    headers=sd.raw_headers, sections=sd.sections,
                )
                continue

            scores = score_color_dataframe(sd.colors)
            labels = score_color_labels(sd.colors)

            key_cols = {}
            for key, aliases in KEY_COLUMNS.items():
                found = _find_column(sd.raw_headers, aliases)
                if found:
                    key_cols[key] = found

            sheets[name] = ProcessedSheet(
                name=name, values=sd.df, colors=sd.colors,
                scores=scores, labels=labels,
                headers=sd.raw_headers, sections=sd.sections,
                key_cols=key_cols,
            )

            all_colors = pd.concat([all_colors, sd.colors], ignore_index=True)

        # 2. Construir master desde "Total Cuentas Gamma" o la hoja mas grande
        total_name = None
        for n in data_sheet_names:
            if "total" in n.lower():
                total_name = n
                break
        if total_name is None:
            total_name = max(data_sheet_names, key=lambda n: len(sheets.get(n, ProcessedSheet(name=n, values=pd.DataFrame(), colors=pd.DataFrame(), scores=pd.DataFrame(), labels=pd.DataFrame())).values))

        total_sheet = sheets.get(total_name)
        if total_sheet is None or total_sheet.values.empty:
            raise ValueError(f"No se encontró hoja maestra: {total_name}")

        master = total_sheet.values.copy()

        # Detectar columnas clave
        col_cliente = total_sheet.key_cols.get("cliente", "CLIENTE")
        col_comercial = total_sheet.key_cols.get("comercial", "Comercial")
        col_zona = total_sheet.key_cols.get("zona", "Zona")
        col_sector = total_sheet.key_cols.get("sector", "SECTOR")
        col_tipo = total_sheet.key_cols.get("tipo_cuenta", "TIPO DE CUENTA (A/B/C)")

        # Normalizar campos
        if col_sector in master.columns:
            master[col_sector] = master[col_sector].apply(_normalize_sector)
        if col_comercial in master.columns:
            master[col_comercial] = master[col_comercial].apply(_normalize_comercial)

        # 3. Extraer heatmap de hojas BDM
        bdm_sheet_names = [n for n in data_sheet_names if n.lower() != total_name.lower()]
        heatmap_df, heatmap_colors = _build_heatmap_matrix_from_bdm_sheets(wb, bdm_sheet_names)

        # Merge heatmap scores into master by client name
        hm_product_cols = []
        hm_service_cols = []

        if not heatmap_df.empty and col_cliente in master.columns:
            master["_cliente_key"] = master[col_cliente].apply(
                lambda x: str(x).strip().upper() if not pd.isna(x) else ""
            )

            # Agrupar scores por cliente (max score si hay duplicados)
            score_cols = [c for c in HEATMAP_PRODUCTS if c in heatmap_df.columns]
            score_svc_cols = [c for c in HEATMAP_SERVICES if c in heatmap_df.columns]
            all_score_cols = score_cols + score_svc_cols

            if all_score_cols:
                # Eliminar columnas conflictivas del master antes del merge
                # (la hoja Total tiene estas columnas con texto, no scores)
                conflicting = [c for c in all_score_cols if c in master.columns]
                if conflicting:
                    master.drop(columns=conflicting, inplace=True)

                agg = heatmap_df.groupby("_cliente_key")[all_score_cols].max().reset_index()
                master = master.merge(agg, on="_cliente_key", how="left")

                # Also merge text values
                text_cols = [f"{c}_text" for c in all_score_cols if f"{c}_text" in heatmap_df.columns]
                if text_cols:
                    # Drop conflicting text cols too
                    conflicting_text = [c for c in text_cols if c in master.columns]
                    if conflicting_text:
                        master.drop(columns=conflicting_text, inplace=True)
                    agg_text = heatmap_df.groupby("_cliente_key")[text_cols].first().reset_index()
                    master = master.merge(agg_text, on="_cliente_key", how="left")

                hm_product_cols = [c for c in score_cols if c in master.columns]
                hm_service_cols = [c for c in score_svc_cols if c in master.columns]

            master.drop(columns=["_cliente_key"], inplace=True, errors="ignore")

        # Fill NaN scores with 0
        for col in hm_product_cols + hm_service_cols:
            if col in master.columns:
                master[col] = master[col].fillna(0).astype(int)

        # Calcular score total heatmap
        if hm_product_cols or hm_service_cols:
            master["score_heatmap"] = master[hm_product_cols + hm_service_cols].sum(axis=1)
            master["cobertura_productos"] = (master[hm_product_cols + hm_service_cols] > 0).sum(axis=1)
        else:
            master["score_heatmap"] = 0
            master["cobertura_productos"] = 0

        # Merge unique colors
        all_unique = detect_unique_colors(all_colors) if not all_colors.empty else {}
        all_unique.update(heatmap_colors)

    finally:
        wb.close()

    return DashboardData(
        master=master,
        sheets=sheets,
        sheet_names=data_sheet_names,
        all_sheet_names=all_sheet_names,
        excel_bytes=excel_bytes,
        unique_colors=all_unique,
        col_cliente=col_cliente,
        col_comercial=col_comercial,
        col_zona=col_zona,
        col_sector=col_sector,
        col_tipo=col_tipo,
        heatmap_product_cols=hm_product_cols,
        heatmap_service_cols=hm_service_cols,
    )


# ---------------------------------------------------------------------------
# Helpers de análisis reutilizables
# ---------------------------------------------------------------------------

def get_heatmap_matrix(data: DashboardData, index_col: str | None = None) -> pd.DataFrame:
    """Retorna matriz pivotada: filas=clientes, cols=productos, valores=scores."""
    idx = index_col or data.col_cliente
    cols = data.heatmap_all_cols
    if not cols or idx not in data.master.columns:
        return pd.DataFrame()

    available = [c for c in cols if c in data.master.columns]
    if not available:
        return pd.DataFrame()

    matrix = data.master[[idx] + available].dropna(subset=[idx]).copy()
    matrix = matrix[matrix[available].sum(axis=1) > 0]
    matrix = matrix.set_index(idx)
    return matrix
